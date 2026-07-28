"""Build the intern API-request guide as a .docx (docs/making-api-requests.docx)."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

doc = Document()

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)


def code_block(text: str):
    """Monospace paragraph with a light grey background."""
    for line in text.rstrip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Pt(12)
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "F2F2F2")
        p.paragraph_format.element.get_or_add_pPr().append(shd)
    doc.add_paragraph()


def para(text: str, bold_parts=()):
    p = doc.add_paragraph()
    for chunk in split_chunks(text, bold_parts):
        run = p.add_run(chunk["text"])
        run.bold = chunk["bold"]
    return p


def split_chunks(text, bold_parts):
    """Split text so listed substrings come out bold."""
    chunks = [{"text": text, "bold": False}]
    for part in bold_parts:
        new = []
        for c in chunks:
            if c["bold"] or part not in c["text"]:
                new.append(c)
                continue
            before, after = c["text"].split(part, 1)
            if before:
                new.append({"text": before, "bold": False})
            new.append({"text": part, "bold": True})
            if after:
                new.append({"text": after, "bold": False})
        chunks = new
    return chunks


title = doc.add_heading("Fetch products from the backend and display them", 0)
title.alignment = WD_ALIGN_PARAGRAPH.LEFT

para(
    "Our backend lives at https://jimmywebsitebackend.vercel.app. "
    "To get products you call:",
)
code_block("GET https://jimmywebsitebackend.vercel.app/api/products?limit=8")

para(
    "No API key is needed — the product catalog is public. Anything you ship in "
    "frontend code is visible to every visitor anyway, so a key there would protect "
    "nothing. Private data like a customer's profile or orders is protected "
    "differently: the customer logs in and every request carries their token.",
    bold_parts=("No API key is needed",),
)

para("Every response from our backend is wrapped like this:")
code_block(
    """{
  "success": true,
  "data": {
    "products": [
      { "id": "...", "name": "...", "price": 25, "currency": "USD", "images": ["..."] }
    ]
  }
}"""
)
para(
    "So after fetching, the products are at json.data.products.",
    bold_parts=("json.data.products",),
)

doc.add_heading("The example — one React component", level=1)
code_block(
    '''"use client";

import { useEffect, useState } from "react";

type Product = {
  id: string;
  name: string;
  price: number;
  currency: string;
  images: string[];
};

export default function ProductList() {
  const [products, setProducts] = useState<Product[]>([]);

  useEffect(() => {
    async function loadProducts() {
      const res = await fetch(
        "https://jimmywebsitebackend.vercel.app/api/products?limit=8"
      );
      const json = await res.json();
      setProducts(json.data.products);
    }
    loadProducts();
  }, []);

  return (
    <ul>
      {products.map((product) => (
        <li key={product.id}>
          <img src={product.images[0]} alt={product.name} width={100} />
          <p>{product.name}</p>
          <p>
            {product.price} {product.currency}
          </p>
        </li>
      ))}
    </ul>
  );
}'''
)

para("That's it:")
steps = [
    "fetch(url) sends the GET request.",
    "res.json() parses the response.",
    "json.data.products is the list — save it in state.",
    ".map() renders one <li> per product.",
]
for s in steps:
    doc.add_paragraph(s, style="List Number")

doc.add_paragraph()
para(
    "Try it: change limit=8 to limit=3, or add &search=shirt to the URL.",
    bold_parts=("Try it:",),
)

OUT = r"C:\Users\jimmy\Desktop\Projects\e-com\frontend-boiler\docs\making-api-requests.docx"
doc.save(OUT)
print(f"Saved {OUT}")
